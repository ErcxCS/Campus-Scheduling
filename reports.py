import os
import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
# Updated imports
from entities import Room, Course
from utils import read_dfs


def build_timetable(courses, rooms, horizon, n_days, solver, start_vars, in_room_vars, seat_vars, exp_path):
    room_ids = [r.id for r in rooms]
    timetable = pd.DataFrame("", index=range(horizon), columns=room_ids)

    for e in courses:
        s_val = solver.Value(start_vars[e.id])
        dur = e.get_duration()
        for t in range(s_val, s_val + dur):
            for r in rooms:
                if solver.Value(in_room_vars[(e.id, r.id)]) == 1:
                    used_seats = solver.Value(seat_vars[(e.id, r.id)])
                    info = f"{e.dep_short} {e.course_code}({used_seats}:{e.year})"
                    if timetable.at[t, r.id]:
                        timetable.at[t, r.id] += "|" + info
                    else:
                        timetable.at[t, r.id] = info

    time_indices = [datetime.time(h + 8, 30).strftime("%H:%M") for h in range(9)]
    day_length = horizon // n_days
    all_days = []
    day_names = [f"Day-{i}" for i in range(1, n_days + 1)]

    for d in range(n_days):
        start_row = d * day_length
        end_row = (d + 1) * day_length
        df_day = timetable.iloc[start_row:end_row, :].copy()
        df_day.index = time_indices
        df_day.columns = [
            f"{rooms[i].room_code}({rooms[i].capacity if rooms[i].is_lab else rooms[i].capacity // 2})"
            for i in range(len(rooms))
        ]
        df_day["Day"] = day_names[d]
        all_days.append((day_names[d], df_day))

    combined = pd.concat([df for _, df in all_days])
    raw_schedule_path = os.path.join(exp_path, "exam_schedule.xlsx")
    combined.to_excel(raw_schedule_path, index=False)
    return all_days

def department_exam_schedule(departments, courses, rooms, horizon, n_days, solver, start_vars, in_room_vars, exp_path):
    day_length = horizon // n_days
    timetables_path = os.path.join(exp_path, "department_schedules")
    os.makedirs(timetables_path, exist_ok=True)
    
    for dep in departments:
        data = []
        for year, year_courses in dep.curriculums.items():
            for course in year_courses:
                start_var = solver.Value(start_vars[course.id])
                nth_day = start_var // day_length + 1
                start_time = start_var % day_length
                
                assigned_rooms = ""
                for r in rooms:
                    if solver.Value(in_room_vars[(course.id, r.id)]) == 1:
                        assigned_rooms += r.room_code + ","
                assigned_rooms = assigned_rooms[:-1]
                data.append({
                    'Year': year,
                    'Course ID': course.course_code,
                    'Day': nth_day,
                    'Starting Slot': start_time,
                    'Rooms': assigned_rooms
                })
        dep_df = pd.DataFrame(data)
        excel_name = dep.name + ".xlsx"
        doc_path = os.path.join(timetables_path, excel_name)
        dep_df.to_excel(doc_path, index=False)

def faculty_exam_schedule(courses, rooms, horizon, n_days, solver, start_vars, in_room_vars, exp_path):
    day_length = horizon // n_days
    data = []
    for course in courses:
        start_var = solver.Value(start_vars[course.id])
        nth_day = start_var // day_length + 1
        start_time = start_var % day_length
        
        assigned_rooms = ""
        total_room_cap = 0
        for r in rooms:
            if solver.Value(in_room_vars[(course.id, r.id)]) == 1:
                assigned_rooms += r.room_code + ","
                total_room_cap += r.capacity if r.is_lab else r.capacity // 2

        assigned_rooms = assigned_rooms[:-1]
        num_rooms_used = len(assigned_rooms.split(","))

        deps = course.get_dep_shorts().split(" ")
        deps = ",".join(deps)
        data.append({
            'Departments': deps,
            'Year': course.year,
            'Course ID': course.course_code,
            'Day': nth_day,
            'Starting Slot': start_time,
            'Assigned Rooms': assigned_rooms,
            'Num Rooms Used': num_rooms_used,
            'Total Room Cap': total_room_cap,
            "Num Students": course.n_students
        })
    fac_df = pd.DataFrame(data)
    excel_name = "faculty_schedule.xlsx"
    doc_path = os.path.join(exp_path, excel_name)
    fac_df.to_excel(doc_path, index=False)

def color_offtimes_black(xlsx_path: str):
    if Room.off_times_dict is None:
        return

    wb = load_workbook(xlsx_path)
    ws = wb["Schedule"]
    black_fill = PatternFill(bgColor="000000", fill_type="solid")

    for i, col in enumerate(ws.iter_cols()):
        off_list = None
        for j, cell in enumerate(col):
            if cell.value is not None:
                value = str(cell.value)
                room_code = value.split('(')[0]
                if room_code in Room.off_times_dict.keys():
                    off_list = Room.off_times_dict[room_code]
                    continue
            if off_list and (j - 1) in off_list:
                cell.fill = black_fill
    wb.save(xlsx_path)

def excelify(dep_list: list, exp_path, room_list, output_filename="exam_schedule.xlsx"):
    schedule_path = os.path.join(exp_path, output_filename)
    combined_df = pd.read_excel(schedule_path)

    cols = combined_df.columns.tolist()
    if "Day" in cols and cols[0] != "Day":
        cols.remove("Day")
        cols = ["Day"] + cols
        combined_df = combined_df[cols]

    beautified_path = os.path.join(exp_path, "beautified_" + output_filename)
    writer = pd.ExcelWriter(beautified_path, engine="xlsxwriter")
    combined_df.to_excel(writer, index=False, sheet_name="Schedule")

    workbook = writer.book
    worksheet = writer.sheets["Schedule"]

    num_columns = len(combined_df.columns)
    worksheet.set_column(1, num_columns - 1, 16)
    worksheet.set_column(0, 0, 12)

    header_rows = 1
    current_day = None
    start_idx = None
    group_index = 0
    yellow_format = workbook.add_format({'bg_color': '#FFFFE0', 'align': 'center', 'valign': 'vcenter'})

    for i, row in combined_df.iterrows():
        day = row["Day"]
        if day != current_day:
            if current_day is not None:
                first_row = header_rows + start_idx
                last_row = header_rows + i - 1
                worksheet.merge_range(first_row, 0, last_row, 0, current_day)
                if group_index % 2 == 0:
                    for r_idx in range(first_row, last_row + 1):
                        worksheet.set_row(r_idx, None, yellow_format)
                group_index += 1
            current_day = day
            start_idx = i
    if current_day is not None:
        first_row = header_rows + start_idx
        last_row = header_rows + i
        worksheet.merge_range(first_row, 0, last_row, 0, current_day)
        if group_index % 2 == 0:
            for r_idx in range(first_row, last_row + 1):
                worksheet.set_row(r_idx, None, yellow_format)

    # Department Coloring
    color_list = ["#A9A9A9", "#800080", "#FFA500", "#FFFF00", "#FF0000", "#008000", "#ADD8E6"]
    dept_colors = {dep.short: color_list[i % len(color_list)] for i, dep in enumerate(dep_list)}
    dept_formats = {k: workbook.add_format({'bg_color': v}) for k, v in dept_colors.items()}
    default_format = workbook.add_format({})
    colored_counts = {k: 0 for k in dept_colors.keys()}

    for col in range(1, num_columns):
        row = header_rows
        while row < header_rows + len(combined_df):
            current_val = combined_df.iloc[row - header_rows, col]
            if pd.isna(current_val):
                current_val = ""
            if current_val != "":
                start_merge = row
                next_row = row + 1
                while (next_row < header_rows + len(combined_df) and
                        combined_df.iloc[next_row - header_rows, col] == current_val):
                    next_row += 1
                end_merge = next_row - 1
                merged_length = end_merge - start_merge + 1

                try:
                    dep_short = current_val.split()[0]
                    cell_format = dept_formats.get(dep_short, default_format)
                except:
                    cell_format = default_format
                    dep_short = None

                if dep_short and cell_format != default_format:
                    colored_counts[dep_short] += merged_length

                if end_merge > start_merge:
                    worksheet.merge_range(start_merge, col, end_merge, col, current_val, cell_format)
                else:
                    worksheet.write(start_merge, col, current_val, cell_format)
                row = next_row
            else:
                row += 1

    writer._save()
    color_offtimes_black(beautified_path)
    print(f"Total mission count: {sum(v // 2 for v in colored_counts.values())}")

def unified_manuel_fac_schedule(dfs: dict, course_list: list):
    # This logic rebuilds a faculty schedule from department files
    data = []
    for course in course_list:
        used_rooms = set()
        date_info = list()
        slot_info = list()
        for dep in course.departments:
            if dep.short not in dfs:
                continue
            dep_df = dfs[dep.short]
            try:
                course_entry = dep_df[dep_df["Course ID"] == course.course_code]
                if course_entry.empty: continue
                rooms_used = course_entry["Rooms Used"].values[0]
            except: continue
            
            used_room_codes = rooms_used.split(",")
            for used_room in used_room_codes:
                try:
                    # Safe to use Room here now
                    used_rooms.add(Room.find_by_code(used_room.strip()))
                except:
                    pass

            date = course_entry["Date"].values[0]
            slot = course_entry["timeslots"].values[0]
            slot_info.append(slot)
            date_info.append(date)
        
        if not date_info: continue

        deps = course.get_dep_shorts().split(" ")
        deps = ",".join(deps)
        data.append({
            'Departments': deps,
            'Year': course.year,
            'Course ID': course.course_code,
            'Course Name': course.course_name,
            'Day': date_info[0],
            'Slot': slot_info[0],
            'Assigned Rooms': [room.room_code for room in used_rooms],
            'Num Rooms Used': len(used_rooms),
            'Total Room Cap': sum([room.capacity if room.is_lab else room.capacity // 2 for room in used_rooms]),
            "Num Students": course.n_students
        })
    return pd.DataFrame(data)

def save_manuel_fac(df, exam):
    df_path = "./data/B24_department_schedules_" + exam
    os.makedirs(df_path, exist_ok=True)
    fac_xslx_path = os.path.join(df_path, "faculty_schedule_manuel_2" + exam + ".xlsx")
    df.to_excel(fac_xslx_path, index=False)

def generate_manuel_fac_(exp_no: int, is_midterm: bool):
    exam_str = "midterms" if is_midterm else "finals"
    # Ensure correct path reading in utils
    schedules_path = "./data/B24_department_schedules_" + exam_str
    if os.path.exists(schedules_path):
        # Uses read_dfs from utils (safe) and Course from entities (safe)
        dfs = read_dfs(schedules_path)
        fac_df_manuel = unified_manuel_fac_schedule(dfs, Course.course_list)
        save_manuel_fac(fac_df_manuel, exam_str)
    else:
        print(f"Warning: Path {schedules_path} not found.")
        print("Skipping manual generation.")