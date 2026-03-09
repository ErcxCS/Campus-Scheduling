import os
import pandas as pd
from openpyxl import load_workbook

def get_off_chunks(slot_list):
    """
    From a list of TimeSlot objects (with .id and .is_off),
    return a list of (start, end) for each consecutive off‐chunk.
    """
    off_chunks = []
    current = []

    for s in sorted(slot_list, key=lambda x: x.id):
        if s.is_off:
            current.append(s.id)
        else:
            if current:
                off_chunks.append((current[0], current[-1] + 1))
                current = []
    if current:
        off_chunks.append((current[0], current[-1] + 1))
    return off_chunks

def find_black_cells(file_path, sheet_name, slots_per_day, week_offset, off_timetable):
    wb = load_workbook(file_path)
    ws = wb[sheet_name]

    for i, col in enumerate(ws.iter_cols()):
        room_code = ""
        for j, cell in enumerate(col):
            cell_color = cell.fill.start_color.index
            if j == 0:
                room_code = cell.value.split('\n')[0]
                if room_code not in off_timetable.keys():
                    off_timetable[room_code] = []
                    continue

            if cell_color == 1:
                time = j - 1
                ts = time - (time // slots_per_day) + week_offset
                off_timetable[room_code].append(ts)

    return off_timetable

def midterm_timetable(slots_per_day: int = 9, num_days: int = 10):
    midterm_tb_path = "./data/bahar_midterm.xlsx"
    horizon = slots_per_day * num_days
    off_timetable = {}
    
    if os.path.exists(midterm_tb_path):
        off_timetable = find_black_cells(midterm_tb_path, 'first', slots_per_day, 0, off_timetable)
        off_timetable = find_black_cells(midterm_tb_path, 'second', slots_per_day, horizon // 2, off_timetable)
    return off_timetable

def read_dfs(schedules_path):
    department_schedules = os.listdir(schedules_path)
    # Simple filtering
    department_schedules = [f for f in department_schedules if len(f.split(" ")) >= 2]
    
    dep_dfs = {}
    for dep in department_schedules:
        dep_word = dep.split(" ")
        dep_code = dep_word[0][0] + dep_word[1][0]
        xlsx_path = os.path.join(schedules_path, dep)
        df = pd.read_excel(xlsx_path, index_col=None, header=0)
        dep_dfs[dep_code] = df
    return dep_dfs

def read_fac_xlsxs(experiment_no: int, exam: str = None, dataset: str = "B24"):
    exp_path = f"./runs/exp{experiment_no}_{exam[:-1]}"
    
    faculty_xlsx = "faculty_schedule.xlsx"
    faculty_path = os.path.join(exp_path, faculty_xlsx)
    fac_df_out = pd.read_excel(faculty_path, index_col=None, header=0)

    df_path = f"./data/B24_department_schedules_{exam}"
    faculty_path = os.path.join(df_path, f"faculty_schedule_manuel_2{exam}.xlsx")
    fac_df_manuel = pd.read_excel(faculty_path, index_col=None, header=0)
    
    return fac_df_manuel, fac_df_out

def compact_columns(df: pd.DataFrame) -> pd.DataFrame:
    compact = {}
    for col in df.columns:
        non_empty = [x for x in df[col].tolist() if x != ""]
        padded = non_empty + [""] * (len(df) - len(non_empty))
        compact[col] = padded

    compact_df = pd.DataFrame(compact)
    mask = (compact_df != "").any(axis=1)
    compact_df = compact_df.loc[mask].reset_index(drop=True)
    compact_df.index = range(1, len(compact_df) + 1)
    return compact_df

def automated_df_rebuild(df: pd.DataFrame, num_days: int) -> pd.DataFrame:
    # Parse Assigned Rooms string → list
    df["RoomsList"] = df["Assigned Rooms"].apply(lambda s: [r.strip() for r in s.split(",")])
    rooms = sorted({room for lst in df["RoomsList"] for room in lst})
    
    TOTAL_ROWS = num_days * 9
    timetable = pd.DataFrame("", index=range(1, TOTAL_ROWS + 1), columns=rooms)

    if df["Starting Slot"].min() == 0:
        df["Starting Slot"] = df["Starting Slot"] + 1

    for _, row in df.iterrows():
        day = int(row["Day"])
        start_slot = int(row["Starting Slot"])
        course_id = row["Course ID"]
        room_list = row["RoomsList"]
        row_index = (day - 1) * 9 + start_slot

        for room in room_list:
            if timetable.at[row_index, room] == "":
                timetable.at[row_index, room] = course_id
            else:
                timetable.at[row_index, room] += " | " + course_id

    mask = (timetable != "").any(axis=1)
    flat = timetable.loc[mask].reset_index(drop=True)
    flat.index = range(1, len(flat) + 1)
    flat = compact_columns(flat)
    return flat